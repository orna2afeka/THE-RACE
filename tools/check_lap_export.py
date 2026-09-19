#!/usr/bin/env python3
"""
check_lap_export.py - the per-lap workbook puts each lap's samples on its own tab
================================================================================
    python tools/check_lap_export.py

Builds throwaway stores whose lap structure is KNOWN sample by sample, writes
the per-lap workbook from them, and reads it back. Also runs against
demo_telemetry.db when that file is present, which it is on the pit laptop and
is not in git.

THE BUG THIS EXISTS FOR, and it is one that cannot be seen by looking at the
output. `calculated_lap` is LapTracker.lap_count -- the number of laps
COMPLETED -- so the samples driven during lap L are tagged L-1, and the figures
for lap L (its time, its energy, its distance) are carried on the rows of lap
L+1, where the car holds them constant. Cut the samples by the obvious rule and
every tab is one lap out: a plausible-looking workbook in which "Lap 40" holds
lap 41's driving, the energy on the tab describes lap 40, and nothing on the
page contradicts anything else. db.py's own comment calls this "a wrong answer
that looks completely plausible", and profile_build.check_lap_alignment exists
because of the same trap on the profile side.

So the alignment is PROVEN here rather than trusted:

  WINDOW      every sample in lap L's window is tagged calculated_lap L-1, on a
              store where that mapping was constructed and is therefore known.
  PARTITION   consecutive laps neither overlap (no sample on two tabs) nor
              leave a hole (no sample lost between them).
  AGREEMENT   the summary on each lap's tab is the same row the Laps index
              shows for that lap -- they come from one builder, so they cannot
              drift.
  GAPS        a window is never allowed to stretch across a break in the data
              and swallow another session's driving.
  REFUSALS    an empty lap range and a range past MAX_LAP_SHEETS are refused,
              not silently truncated.
  UNTOUCHED   write_xlsx still produces the Data/Laps/Charts workbook it always
              did; the per-lap export is beside it, not instead of it.

Run it with the window built from lap L instead of L-1 and WINDOW fails.
"""

import os
import sys
import tempfile
import warnings

warnings.filterwarnings("ignore")

_ROOT = os.environ.get("SOLARRACE_ROOT") or os.path.dirname(
    os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, _ROOT)
sys.path.insert(0, os.path.join(_ROOT, "Pit_Dashboard"))

import db                                                        # noqa: E402
import export                                                    # noqa: E402

FAILED = []

PUSH_S = 0.5            # the car's telemetry interval
LAP_S = 200.0           # every synthetic lap takes exactly this long
PER_LAP = int(LAP_S / PUSH_S)


def check(label, ok, detail=""):
    print("  %-54s %s" % (label, "OK  " if ok else "FAIL"))
    if detail:
        print("       " + detail)
    if not ok:
        FAILED.append("%s - %s" % (label, detail))
    return ok


def make_store(path, laps=6, gap_before=None):
    """A store of `laps` driven laps, tagged the way the car tags them.

    Returns {lap number: [device_ts, ...]} for the samples DRIVEN in that lap,
    which is the ground truth every window is measured against.

    The tagging is the car's, not a convenience: while lap L is being driven
    the car reports calculated_lap = L-1 (that many are finished) and carries
    the PREVIOUS lap's time, energy and distance in the last_lap_* columns. So
    the figures for lap L appear only once lap L+1 is under way -- which is why
    the last lap driven is not a completed lap.

    `gap_before` drops an hour of silence in front of that lap number, so the
    window rule can be checked against a break in the data.
    """
    conn = db.get_conn(path)
    db.init_db(conn)
    names = ["device_id", "rtdb_key", "device_ts", "ingested_ts",
             "calculated_lap", "lap_seq", "lap_distance_m", "lap_source",
             "last_lap_number", "last_lap_time_s", "last_lap_energy",
             "last_lap_regen_energy", "last_lap_distance_m", "last_lap_kind",
             "mms_vehicle_speed_kmh", "odometer_m"]
    sql = "INSERT INTO telemetry (%s) VALUES (%s)" % (
        ",".join(names), ",".join("?" * len(names)))

    truth, t, k = {}, 1789560000.0, 0
    for lap in range(1, laps + 1):
        if gap_before == lap:
            t += 3600.0
        truth[lap] = []
        done = lap - 1                       # laps COMPLETED while lap is driven
        for i in range(PER_LAP):
            ts = t + i * PUSH_S
            truth[lap].append(ts)
            conn.execute(sql, [
                "solarcar", "-K%07d" % k, ts, ts,
                float(done), float(done), i * (4000.0 / PER_LAP), "gps",
                # The lap just finished: absent for the very first lap.
                float(done) if done else None,
                LAP_S if done else None,
                100.0 + done if done else None,
                5.0 if done else None,
                4000.0 if done else None,
                "flying" if done else None,
                55.0, 4000.0 * done + i * (4000.0 / PER_LAP),
            ])
            k += 1
        t += LAP_S
    conn.commit()
    conn.close()
    return truth


def tagged_laps(conn, t0, t1):
    """The distinct calculated_lap values of the samples in [t0, t1)."""
    return {r[0] for r in conn.execute(
        "SELECT DISTINCT calculated_lap FROM telemetry "
        "WHERE device_id='solarcar' AND device_ts >= ? AND device_ts < ?",
        (t0, t1))}


def windows_for(conn):
    """[(lap row, (t0, t1) or None)] for every completed lap in a store."""
    laps = db.fetch_laps(conn)
    return laps, [(l, export._lap_window(laps, i)) for i, l in enumerate(laps)]


def check_alignment(tmp):
    path = os.path.join(tmp, "laps.db")
    truth = make_store(path, laps=6)
    conn = db.get_conn(path)
    laps, pairs = windows_for(conn)

    nums = [l["lap"] for l in laps]
    check("the store reports the laps that finished",
          nums == [1, 2, 3, 4, 5],
          "laps %s (lap 6 is still being driven, so it has no figures yet)" % nums)

    off = []
    for lap, win in pairs:
        if not win:
            off.append("lap %s has no window" % lap["lap"])
            continue
        got = tagged_laps(conn, *win)
        want = {float(lap["lap"] - 1)}
        if got != want:
            off.append("lap %s holds calculated_lap %s, wanted %s"
                       % (lap["lap"], sorted(got), sorted(want)))
    check("WINDOW: every sample is tagged one below its lap",
          not off, "; ".join(off) or
          "laps %s each hold only calculated_lap L-1" % nums)

    counts, bad = [], []
    # Windowless laps are already reported above; skipping them here keeps a
    # failing run printing its remaining checks instead of ending in a
    # traceback.
    for lap, win in [(l, w) for l, w in pairs if w]:
        n = len(db.fetch_samples(conn, start_ts=win[0], end_ts=win[1] - 1e-6))
        counts.append(n)
        if n != len(truth[lap["lap"]]):
            bad.append("lap %s: %d samples, drove %d"
                       % (lap["lap"], n, len(truth[lap["lap"]])))
    check("every sample the lap was driven with, and no other",
          not bad, "; ".join(bad) or
          "%d samples per lap, %.0f s at %.1f s" % (counts[0], LAP_S, PUSH_S))

    holes = []
    solid = [(l, w) for l, w in pairs if w]
    for i in range(1, len(solid)):
        prev_end, start = solid[i - 1][1][1], solid[i][1][0]
        if abs(prev_end - start) > 1e-6:
            holes.append("lap %s ends %.3f, lap %s starts %.3f"
                         % (solid[i - 1][0]["lap"], prev_end,
                            solid[i][0]["lap"], start))
    check("PARTITION: laps meet exactly, no overlap and no hole",
          not holes, "; ".join(holes) or
          "%d consecutive boundaries share an instant" % (len(solid) - 1))
    conn.close()


def check_gap(tmp):
    """A break in the data must not be swallowed into the following lap."""
    path = os.path.join(tmp, "gap.db")
    make_store(path, laps=6, gap_before=4)
    conn = db.get_conn(path)
    laps, pairs = windows_for(conn)
    by_lap = {l["lap"]: w for l, w in pairs}
    win = by_lap.get(4)
    span = (win[1] - win[0]) if win else None
    check("GAPS: an hour of silence is not counted as lap time",
          span is not None and span <= LAP_S * export._LAP_WINDOW_SLACK,
          "lap 4 spans %.0f s (the lap took %.0f s; an hour was missing "
          "before it)" % (span or -1, LAP_S))
    # And its samples are still the right ones.
    check("       and the lap still holds its own samples",
          win is not None and tagged_laps(conn, *win) == {3.0},
          "calculated_lap %s" % sorted(tagged_laps(conn, *win)) if win else "")
    conn.close()


def check_workbook(tmp):
    from openpyxl import load_workbook
    path = os.path.join(tmp, "laps.db")
    conn = db.get_conn(path)
    out = os.path.join(tmp, "perlap.xlsx")
    nlaps, nrows = export.write_laps_xlsx(out, first_lap=2, last_lap=4, conn=conn)
    wb = load_workbook(out)

    check("the index comes first, then a tab per lap",
          wb.sheetnames == ["Laps", "Lap 2", "Lap 3", "Lap 4"],
          "sheets %s" % wb.sheetnames)
    check("it returns what it wrote",
          (nlaps, nrows) == (3, 3 * PER_LAP),
          "%d laps, %d rows" % (nlaps, nrows))

    # AGREEMENT: the tab's summary is the index's row for the same lap.
    index = wb["Laps"]
    idx_rows = {r[0].value: [c.value for c in r[:len(export._LAP_HEADERS)]]
                for r in index.iter_rows(min_row=2) if isinstance(r[0].value, int)}
    drift = []
    for lap in (2, 3, 4):
        tab = [c.value for c in wb["Lap %d" % lap][2][:len(export._LAP_HEADERS)]]
        if tab != idx_rows.get(lap):
            drift.append("lap %d: tab %s vs index %s"
                         % (lap, tab[:5], (idx_rows.get(lap) or [])[:5]))
    check("AGREEMENT: each tab's summary is its row in the index",
          not drift, "; ".join(drift) or "3 tabs match the index row for row")

    ws = wb["Lap 3"]
    check("the samples start under both header rows",
          ws.freeze_panes == "A5" and ws.max_row == 4 + PER_LAP,
          "freeze %s, %d rows" % (ws.freeze_panes, ws.max_row))
    check("the raw lap column is not labelled like the tab",
          "Lap" not in [c.value for c in ws[4]],
          "row 4 headers say %s" % [c.value for c in ws[4]][:3])
    conn.close()


def check_refusals(tmp):
    path = os.path.join(tmp, "laps.db")
    conn = db.get_conn(path)
    out = os.path.join(tmp, "refused.xlsx")

    def refuses(**kw):
        limit = kw.pop("limit", None)
        was = export.MAX_LAP_SHEETS
        if limit:
            export.MAX_LAP_SHEETS = limit
        try:
            export.write_laps_xlsx(out, conn=conn, **kw)
            return None
        except ValueError as e:
            return str(e)
        finally:
            export.MAX_LAP_SHEETS = was

    why = refuses(first_lap=900, last_lap=999)
    check("REFUSALS: a range with no lap in it is refused",
          why is not None, why or "it wrote a workbook with no laps")
    why = refuses(limit=2)
    check("            and one past the sheet limit",
          why is not None and "limit" in why, why or "it wrote every sheet")
    conn.close()


def check_duplicate_numbers(tmp):
    """The pit can set the lap number, so two laps CAN share one."""
    from openpyxl import Workbook
    wb = Workbook()
    # Named and created one at a time, the way write_laps_xlsx does it: the
    # name is chosen against the sheets that already exist.
    titles = []
    for _ in range(3):
        titles.append(export._lap_sheet_title(wb, {"lap": 7}))
        wb.create_sheet(titles[-1])
    check("a repeated lap number still gets its own tab",
          titles == ["Lap 7", "Lap 7 (2)", "Lap 7 (3)"], "titles %s" % titles)


def check_existing_export(tmp):
    path = os.path.join(tmp, "laps.db")
    conn = db.get_conn(path)
    out = os.path.join(tmp, "whole.xlsx")
    n = export.write_xlsx(out, conn=conn)
    from openpyxl import load_workbook
    sheets = load_workbook(out).sheetnames
    check("UNTOUCHED: the time-ranged workbook is as it was",
          n == 6 * PER_LAP and sheets[0] == "Data" and "Laps" in sheets,
          "%d rows, sheets %s" % (n, sheets))
    conn.close()


def check_double_cuts():
    """A double cut's phantom is left out -- and NOTHING else is."""
    def lap(n, m, s, source="manual", flags=("distance_suspect",)):
        return {"lap": n, "distance_m": m, "lap_time_s": s,
                "lap_source": source, "flags": list(flags)}

    phantom = lap(59, 140.0, 9.1)                 # 2026-09-19 18:03:23, verbatim
    kept = [
        lap(58, 4000.0, 306.6, "odometer", ("virtual_end",)),
        lap(60, 140.0, 9.1, "manual", ()),            # short, but the car did not flag it
        lap(61, 140.0, 9.1, "odometer"),              # flagged, but not cut by hand
        lap(62, 3920.0, 1335.0),                      # flagged and manual, but a real lap
        lap(63, 300.0, 400.0),                        # short, but it took minutes: a stop
        {"lap": 64, "distance_m": None, "lap_time_s": None,
         "lap_source": None, "flags": []},            # an old car: nothing to judge by
    ]
    out = db._drop_double_cuts([phantom] + kept)
    check("PHANTOM: a 9 s, 140 m hand cut the car flagged is left out",
          phantom not in out, "lap 59 of 140 m in 9.1 s")
    check("         and every other lap stays listed",
          out == kept, "%d of %d kept" % (len(out), len(kept)))


def check_real_store():
    """The demo store, when the laptop has one. Not in git; skipped elsewhere."""
    real = os.path.join(_ROOT, "demo_telemetry.db")
    if not os.path.exists(real):
        print("  (demo_telemetry.db not present - skipping the real store)")
        return
    import sqlite3
    conn = sqlite3.connect("file:%s?mode=ro" % real.replace("\\", "/"), uri=True)
    conn.row_factory = sqlite3.Row
    laps, pairs = windows_for(conn)
    off = [l["lap"] for l, w in pairs
           if w and tagged_laps(conn, *w) != {float(l["lap"] - 1)}]
    check("the demo store aligns the same way",
          not off and len(pairs) > 1,
          "%d laps, %d misaligned" % (len(pairs), len(off)))
    conn.close()


def main():
    print("\nthe per-lap workbook, against a store whose laps are known\n")
    tmp = tempfile.mkdtemp(prefix="lapexport_")
    check_alignment(tmp)
    check_gap(tmp)
    check_workbook(tmp)
    check_refusals(tmp)
    check_duplicate_numbers(tmp)
    check_existing_export(tmp)
    check_double_cuts()
    check_real_store()
    if FAILED:
        print("\n%d FAILED:" % len(FAILED))
        for f in FAILED:
            print("  - %s" % f)
        return 1
    print("\nAll per-lap export checks passed.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
